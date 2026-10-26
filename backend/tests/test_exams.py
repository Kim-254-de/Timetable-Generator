"""
The exam timetable: departments publish papers in Smart Attendance, the Director
sets the exam period, imports, generates, shares the draft for invigilators,
and publishes.
"""

from __future__ import annotations

import uuid
from collections import defaultdict
from io import BytesIO
from typing import Any

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import ExamPublication
from tests.conftest import sign_in
from tests.fake_teaching import FakeTeaching, course, paper


def ok(response: Any) -> Any:
    assert response.status_code == 200, response.text
    return response.json()["data"]


def setup_exams(client: TestClient, teaching: FakeTeaching) -> str:
    term = ok(
        client.post(
            "/api/terms",
            json={
                "academic_year": "2026/2027",
                "semester": 1,
                "name": "Semester I 2026/2027",
                "start_date": "2026-09-07",
                "end_date": "2026-12-18",
            },
        )
    )
    rows = [
        {"code": "ADM 1", "building": "ADM", "capacity": "100", "room_type": "Lecture hall"},
        {"code": "STB 1", "building": "STB", "capacity": "200", "room_type": "Lecture hall"},
        {"code": "STB 2", "building": "STB", "capacity": "200", "room_type": "Lecture hall"},
        {"code": "UTC 1", "building": "UTC", "capacity": "300", "room_type": "Lecture hall"},
    ]
    ok(client.post("/api/rooms/bulk", json={"rows": rows}))
    ok(
        client.patch(
            f"/api/terms/{term['id']}/exam-period",
            json={"exam_start": "2026-11-23", "exam_end": "2026-11-27", "exam_days_off": ["2026-11-26"]},
        )
    )
    teaching.exam_courses = [
        course("c-cs1", "CS", "BSC.CS", 120),
        course("c-cs3", "CS", "BSC.CS", 20, year=3),
        {**course("c-odel", "CS", "BSC.CS", 30), "sub_class": "ODEL", "label": "BSC.CS Y1S1 (ODEL)"},
        course("c-ed1", "ED", "BED.SCI", 80),
        course("c-ed2", "ED", "BED.ARTS", 300, year=2),
    ]
    teaching.exam_published = [
        {
            "code": "CS",
            "name": "Computer Science",
            "published_at": "2026-10-23T08:00:00+00:00",
            "published_by": "A Cod",
            "papers": [
                paper(
                    "p-shared",
                    ["BITE 131", "COSC 130"],
                    [("c-cs1", 120, "regular"), ("c-odel", 30, "odel"), ("c-ed1", 80, "regular")],
                ),
                paper("p-200", ["COSC 200"], [("c-cs1", 120, "regular")]),
                paper("p-301", ["COSC 301"], [("c-cs3", 20, "regular")]),
                paper("p-epsc", ["EPSC 121"], [("c-ed2", 300, "regular")]),
            ],
            "practicals": [
                {
                    "id": "pr-1",
                    "codes": ["COSC 200"],
                    "title": "Data",
                    "date": "2026-11-20",
                    "start_time": "09:00",
                    "end_time": "11:00",
                }
            ],
        }
    ]
    teaching.exam_not_published = [{"code": "ED", "name": "Education", "reopened_at": None}]
    return str(term["id"])


def test_the_exam_period_and_exam_seats(client, teaching):
    sign_in(client)
    term_id = setup_exams(client, teaching)
    period = ok(client.get(f"/api/terms/{term_id}/exams"))["period"]
    assert period["days"] == ["2026-11-23", "2026-11-24", "2026-11-25", "2026-11-27"]  # Thursday off
    assert period["slots"] == 12 and period["exam_seats"] == 90 + 180 + 180 + 270  # 0.9 x capacity
    rooms = {r["code"]: r for r in ok(client.get("/api/rooms"))}
    ok(client.patch(f"/api/rooms/{rooms['UTC 1']['id']}", json={"exam_seats": 300}))
    ok(client.patch(f"/api/rooms/{rooms['ADM 1']['id']}", json={"for_exams": False}))
    assert ok(client.get(f"/api/terms/{term_id}/exams"))["period"]["exam_seats"] == 180 + 180 + 300
    assert rooms["UTC 1"]["block"] == "UTC"


def test_import_generate_share_and_publish(client, teaching, smarttt, db: Session):
    sign_in(client)
    term_id = setup_exams(client, teaching)
    status = ok(client.get(f"/api/terms/{term_id}/exam-departments"))
    assert [d["code"] for d in status["published"]] == ["CS"] and status["not_published"][0]["code"] == "ED"
    result = ok(client.post(f"/api/terms/{term_id}/exams/import"))
    assert sorted(result["papers"]["added"]) == ["BITE 131/COSC 130", "COSC 200", "COSC 301", "EPSC 121"]
    exams = ok(client.get(f"/api/terms/{term_id}/exams"))
    assert exams["readiness"] == [] and exams["practicals"][0]["date"] == "2026-11-20"

    run = ok(client.post("/api/exam-runs", json={"term_id": term_id, "time_limit_s": 20}))
    assert run["status"] == "succeeded", run
    grid = ok(client.get(f"/api/exam-drafts/{run['draft_id']}"))
    assert grid["draft"]["violations"] == [] and grid["draft"]["placed"] == 4
    papers = {p["codes"]: p for p in grid["papers"]}

    # A class row sits at most 2 papers a day, never two at once (BSC.CS Y1S1 sits two papers).
    by_day: dict[str, list[int]] = defaultdict(list)
    for codes in ("BITE 131/COSC 130", "COSC 200"):
        by_day[papers[codes]["day"]].append(papers[codes]["session"])
    assert all(len(set(s)) == len(s) <= 2 for s in by_day.values())
    # The ODEL class row sits in its own room; the big class row (300 > every room) is split over two.
    shared = papers["BITE 131/COSC 130"]["seating"]
    odel_rooms = {s["room"] for s in shared if s["odel"]}
    assert odel_rooms and odel_rooms.isdisjoint({s["room"] for s in shared if not s["odel"]})
    epsc = papers["EPSC 121"]["seating"]
    assert len(epsc) == 2 and {(s["part"], s["parts"]) for s in epsc} == {(1, 2), (2, 2)}
    # Practicals are shown on their own date, not moved into the period.
    assert grid["practicals"][0]["date"] == "2026-11-20"

    # Move COSC 301 beside COSC 200: the small paper shares a room within its seats.
    draft_id = run["draft_id"]
    target = papers["COSC 200"]
    ok(
        client.patch(
            f"/api/exam-drafts/{draft_id}/papers/{papers['COSC 301']['id']}",
            json={"day": target["day"], "session": target["session"]},
        )
    )
    grid = ok(client.get(f"/api/exam-drafts/{draft_id}"))
    papers = {p["codes"]: p for p in grid["papers"]}
    assert {s["room"] for s in papers["COSC 301"]["seating"]} == {s["room"] for s in papers["COSC 200"]["seating"]}
    # A move that breaks a rule is refused: the shared paper into COSC 200's session (BSC.CS Y1S1 sits both).
    clash = client.patch(
        f"/api/exam-drafts/{draft_id}/papers/{papers['BITE 131/COSC 130']['id']}",
        json={"day": target["day"], "session": target["session"]},
    )
    assert clash.status_code == 409 and clash.json()["error"]["code"] == "BREAKS_RULE"

    # Publishing needs the draft shared, then an invigilator for every paper.
    assert client.post(f"/api/exam-drafts/{draft_id}/publish").json()["error"]["code"] == "NOT_SHARED"
    ok(client.post(f"/api/exam-drafts/{draft_id}/share"))
    shared_view = ok(
        client.get(
            f"/api/integrations/attendance/terms/{term_id}/exam-draft", headers={"X-API-Key": "test-attendance-key"}
        )
    )
    assert {p["paper_ref"] for p in shared_view["papers"]} == {"p-shared", "p-200", "p-301", "p-epsc"}
    assert all(p["rooms"] and p["start_time"] for p in shared_view["papers"])
    # Codes as a list, the way Smart Attendance's invigilator page reads them.
    shared_paper = next(p for p in shared_view["papers"] if p["paper_ref"] == "p-shared")
    assert shared_paper["codes"] == ["BITE 131", "COSC 130"] and shared_paper["sitting"] == "regular"
    refused = client.post(f"/api/exam-drafts/{draft_id}/publish")
    assert refused.status_code == 409 and refused.json()["error"]["code"] == "NO_INVIGILATOR"
    teaching.invigilators = {
        ref: {"": {"staff_number": f"STF/{n}", "name": f"Dr. Invigilator {n}"}}
        for n, ref in enumerate(["p-shared", "p-200", "p-301", "p-epsc"])
    }
    published = ok(client.post(f"/api/exam-drafts/{draft_id}/publish"))
    assert published["version"] == 1 and published["smarttt"]["sent"] is True

    publication = db.scalar(select(ExamPublication).where(ExamPublication.draft_id == uuid.UUID(draft_id)))
    assert publication is not None
    by_key = {p["key"]: p for p in publication.payload["papers"]}
    seat = by_key["p-200"]["seating"][0]
    assert seat["room_code"] and seat["students"] == 120 and seat["invigilator_staff_number"] == "STF/1"
    assert by_key["p-shared"]["codes"] == ["BITE 131", "COSC 130"] and by_key["p-shared"]["start_time"]
    # Each class row's programme, year and stream, so SMARTTT can show students their own room.
    assert (seat["program_code"], seat["year_of_study"], seat["study_semester"]) == ("BSC.CS", 1, 1)
    odel = next(s for s in by_key["p-shared"]["seating"] if s["class_row"] == "BSC.CS Y1S1 (ODEL)")
    assert odel["stream"] == "ODEL"

    # SMARTTT received it, with the practicals.
    sent = smarttt.exam_payloads[-1]
    assert sent["version"] == 1 and sent["based_on_version"] is None
    assert sent["practicals"][0]["codes"] == ["COSC 200"] and sent["practicals"][0]["date"] == "2026-11-20"
    assert ok(client.get(f"/api/terms/{term_id}/exams"))["publications"][0]["sent_to_smarttt"] is True

    # Downloads: the timetable, room sheets and roster.
    timetable = client.get(f"/api/exam-drafts/{draft_id}/export/timetable.docx")
    assert timetable.status_code == 200
    doc = Document(BytesIO(timetable.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "DIRECTORATE OF EXAMINATION & TIMETABLING" in text and "VERSION 1" in text
    cells = " ".join(c.text for row in doc.tables[0].rows for c in row.cells)
    assert "DATE/DAY" in cells and "COSC 200 →" in cells
    for kind in ("timetable", "rooms", "roster"):
        assert client.get(f"/api/exam-drafts/{draft_id}/export/{kind}.pdf").content[:4] == b"%PDF"
    roster = Document(BytesIO(client.get(f"/api/exam-drafts/{draft_id}/export/roster.docx").content))
    assert "Dr. Invigilator 1" in " ".join(c.text for t in roster.tables for row in t.rows for c in row.cells)


def test_generating_needs_the_period_rooms_and_papers(client, teaching):
    sign_in(client)
    term = ok(client.post("/api/terms", json={"academic_year": "2026/2027", "semester": 1}))
    ok(client.post("/api/rooms/bulk", json={"rows": [{"code": "A", "building": "A", "capacity": "50"}]}))
    refused = client.post("/api/exam-runs", json={"term_id": term["id"]})
    assert refused.status_code == 409
    codes = {p["code"] for p in refused.json()["error"]["details"]["problems"]}
    assert codes == {"no_period", "no_papers"}


def _generate(client: TestClient, term_id: str) -> str:
    run = ok(client.post("/api/exam-runs", json={"term_id": term_id, "time_limit_s": 20}))
    assert run["status"] == "succeeded", run
    return str(run["draft_id"])


def test_an_exam_version_smarttt_missed_is_sent_later(client, teaching, smarttt):
    sign_in(client)
    term_id = setup_exams(client, teaching)
    ok(client.post(f"/api/terms/{term_id}/exams/import"))
    draft_id = _generate(client, term_id)
    ok(client.post(f"/api/exam-drafts/{draft_id}/share"))
    teaching.invigilators = {
        ref: {"": {"staff_number": f"STF/{n}", "name": f"Dr. Invigilator {n}"}}
        for n, ref in enumerate(["p-shared", "p-200", "p-301", "p-epsc"])
    }
    smarttt.down = True
    published = ok(client.post(f"/api/exam-drafts/{draft_id}/publish"))
    assert published["version"] == 1 and published["smarttt"]["sent"] is False
    (version,) = ok(client.get(f"/api/terms/{term_id}/exams"))["publications"]
    assert version["sent_to_smarttt"] is False and version["can_send"] is True

    down = client.post(f"/api/terms/{term_id}/exam-publications/1/send")
    assert down.status_code == 503
    smarttt.down = False
    sent = ok(client.post(f"/api/terms/{term_id}/exam-publications/1/send"))
    assert sent["sent_to_smarttt"] is True and sent["can_send"] is False
    assert smarttt.exam_publications[1]["paper_count"] == 4
    again = client.post(f"/api/terms/{term_id}/exam-publications/1/send")
    assert again.json()["error"]["code"] == "ALREADY_SENT"


def test_an_odel_paper_of_the_same_unit_is_its_own_paper(client, teaching, smarttt):
    sign_in(client)
    term_id = setup_exams(client, teaching)
    # COSC 200's ODEL class sits a separate paper, as Smart Attendance sends it.
    teaching.exam_published[0]["papers"].append(
        {**paper("p-200-odel", ["COSC 200"], [("c-odel", 30, "odel")]), "sitting": "odel"}
    )
    result = ok(client.post(f"/api/terms/{term_id}/exams/import"))
    assert "COSC 200 (ODEL)" in result["papers"]["added"] and "COSC 200" in result["papers"]["added"]
    draft_id = _generate(client, term_id)
    papers = {p["codes"]: p for p in ok(client.get(f"/api/exam-drafts/{draft_id}"))["papers"]}
    assert papers["COSC 200 (ODEL)"]["sitting"] == "odel" and papers["COSC 200"]["sitting"] == "regular"
    assert {s["row"] for s in papers["COSC 200 (ODEL)"]["seating"]} == {"BSC.CS Y1S1 (ODEL)"}

    ok(client.post(f"/api/exam-drafts/{draft_id}/share"))
    teaching.invigilators = {
        ref: {"": {"staff_number": f"STF/{n}", "name": f"Dr. Invigilator {n}"}}
        for n, ref in enumerate(["p-shared", "p-200", "p-301", "p-epsc", "p-200-odel"])
    }
    assert ok(client.post(f"/api/exam-drafts/{draft_id}/publish"))["smarttt"]["sent"] is True
    sent = {p["key"]: p for p in smarttt.exam_payloads[-1]["papers"]}
    # SMARTTT gets the codes as they are, and the sitting beside them.
    assert (sent["p-200-odel"]["codes"], sent["p-200-odel"]["sitting"]) == (["COSC 200"], "odel")
    assert sent["p-200"]["sitting"] == "regular"


def test_the_timetable_pdf_holds_a_session_busier_than_a_page():
    """At full size (899 papers) one session had 86 papers: its day must not overflow the page."""
    from types import SimpleNamespace

    from app.services import exam_export

    day = {"date": "2026-12-01", "label": "TUE 01/12/2026"}
    papers = [
        {"codes": f"COSC {100 + i}", "day": day["date"], "session": 0, "seating": [{"room": f"UTC {i % 14}"}]}
        for i in range(150)
    ]
    rec = {"days": [day], "sessions": [{"start": "08:30", "end": "10:30"}], "papers": papers, "practicals": []}
    term = SimpleNamespace(
        academic_year="2026/2027", semester=1, name="Semester I", prepared_by="", prepared_by_title=""
    )
    pdf = exam_export.timetable_pdf(term, rec, 1, "Director")
    assert pdf[:4] == b"%PDF"
    rows = exam_export._fit_pages(exam_export.timetable_table(rec)[1])
    assert [r[0] for r in rows] == ["TUE 01/12/2026"] + ["TUE 01/12/2026 (cont.)"] * 3
