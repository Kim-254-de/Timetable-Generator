"""
The Directorate enters everything in the Generator, with no SMARTTT data:
a term, lecture halls, departments, class rows and units; then it generates,
publishes and downloads the timetable as PDF, Word and Excel.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from fastapi.testclient import TestClient
from sqlalchemy import select

from app.models import UnitOffering
from tests.conftest import departments, sign_in
from tests.fake_smarttt import CS, EDU


def ok(response: Any) -> Any:
    assert response.status_code == 200, response.text
    return response.json()["data"]


def enter_term(client: TestClient) -> dict[str, Any]:
    """The client's example: COSC 100 for one class row, COSC 103 split into GR A/B/C."""
    term = ok(
        client.post(
            "/api/terms",
            json={
                "academic_year": "2026/2027",
                "semester": 1,
                "name": "September-December 2026",
                "version_label": "VERSION_1 DATED 19/09/2026",
                "use_from": "2026-09-21",
                "prepared_by": "Dr. Alex Mutegi Rwanda",
            },
        )
    )
    assert term["smarttt_term_id"] is None and term["default_session_length"] == 3

    result = ok(
        client.post(
            "/api/rooms/bulk",
            json={
                "rows": [
                    {"code": "UTC 1", "building": "UTC", "capacity": "200", "room_type": "Lecture hall"},
                    {"code": "STB 1", "building": "STB", "capacity": "120", "room_type": "Lecture hall"},
                    {"code": "STB 2", "building": "STB", "capacity": "120", "room_type": "lecture hall"},
                    {"code": "STB 3", "building": "STB", "capacity": "120", "room_type": "Hall"},
                    {"code": "LAB 1", "building": "STB", "capacity": "40", "room_type": "Computer lab"},
                    {"code": "BAD", "building": "X", "capacity": "lots"},
                ]
            },
        )
    )
    assert result["created"] == 5 and result["error_count"] == 1 and result["errors"][0]["line"] == 6

    # Departments come from Smart Attendance; the Generator sets its own codes and lecturer counts.
    found = departments(
        client, {"Computer Science": {"code": "DCOMP", "lecturer_count": 6}, "Education": {"code": "DEDU"}}
    )
    cs, edu = found["Computer Science"], found["Education"]
    assert (cs["id"], edu["id"]) == (CS, EDU) and cs["from_attendance"] and cs["attendance_lecturer_count"] == 5

    rows = ok(
        client.post(
            f"/api/terms/{term['id']}/class-rows/bulk",
            json={
                "rows": [
                    {
                        "programme": "BSC.COMP SCI",
                        "year": "1",
                        "semester": "1",
                        "department": "DCOMP",
                        "students": "90",
                    },
                    {
                        "programme": "BED.ENG/LIT",
                        "year": "1",
                        "semester": "1",
                        "sub_class": "1",
                        "department": "DEDU",
                        "students": "110",
                    },
                    {
                        "programme": "BED.KISW/RELI",
                        "year": "1",
                        "semester": "1",
                        "department": "DEDU",
                        "students": "100",
                    },
                    {
                        "programme": "BED.MATH/BIO",
                        "year": "1",
                        "semester": "1",
                        "department": "DEDU",
                        "students": "105",
                    },
                ]
            },
        )
    )
    assert rows["created"] == 4, rows
    labels = {r["label"]: r["id"] for r in ok(client.get(f"/api/terms/{term['id']}/class-rows"))}
    assert "BED.ENG/LIT Y1S1 (1)" in labels

    cosc100 = ok(
        client.post(
            f"/api/terms/{term['id']}/offerings",
            json={
                "unit_code": "COSC 100",
                "unit_name": "Introduction to Computing",
                "department_id": cs["id"],
                "class_row_ids": [labels["BSC.COMP SCI Y1S1"]],
            },
        )
    )
    # Hours default to the term's 3-hour block; a unit with sized class rows starts as one whole class.
    assert cosc100["weekly_hours"] == 3 and cosc100["session_length"] == 3
    assert [g["label"] for g in cosc100["groups"]] == [""]

    education = [labels["BED.ENG/LIT Y1S1 (1)"], labels["BED.KISW/RELI Y1S1"], labels["BED.MATH/BIO Y1S1"]]
    cosc103 = ok(
        client.post(
            f"/api/terms/{term['id']}/offerings",
            json={
                "unit_code": "COSC 103",
                "unit_name": "Introduction to ICT Skills",
                "department_id": cs["id"],
                "class_row_ids": education,
            },
        )
    )
    whole = cosc103["groups"][0]
    ok(client.delete(f"/api/groups/{whole['id']}"))
    for label, row_id, students in zip(("GR A", "GR B", "GR C"), education, (110, 100, 105), strict=True):
        ok(
            client.post(
                f"/api/offerings/{cosc103['id']}/groups",
                json={"label": label, "members": [{"class_row_id": row_id, "students": students}]},
            )
        )
    return term


def test_director_enters_everything_and_downloads(client, smarttt, db):
    sign_in(client)
    term = enter_term(client)
    overview = ok(client.get(f"/api/terms/{term['id']}"))
    assert overview["readiness"]["ready"] is True, overview["readiness"]["problems"]

    run = ok(client.post("/api/runs", json={"term_id": term["id"], "time_limit_s": 20}))
    assert run["status"] == "succeeded", run["failure"]
    draft_id = run["draft_id"]
    grid = ok(client.get(f"/api/drafts/{draft_id}?view=row"))
    assert grid["draft"]["violation_count"] == 0
    groups = {c["group"] for e in grid["entities"] for c in e["cells"] if c["unit_code"] == "COSC 103"}
    assert groups == {"GR A", "GR B", "GR C"}

    # Published without SMARTTT: the approval stands, flagged as not sent.
    published = ok(client.post(f"/api/drafts/{draft_id}/publish", json={}))
    assert published["version"] == 1 and published["smarttt"]["sent"] is False
    assert smarttt.payloads == []
    version = ok(client.get(f"/api/terms/{term['id']}"))["publications"][0]
    assert version["sent_to_smarttt"] is False and version["can_send"] is False
    assert "start and end dates" in version["not_sent_reason"]

    # With its dates, SMARTTT can create the term: the version is sent and the
    # term (and the units SMARTTT matched or created) linked to SMARTTT's ids.
    ok(client.patch(f"/api/terms/{term['id']}", json={"start_date": "2026-09-21", "end_date": "2026-12-18"}))
    sent = ok(client.post(f"/api/terms/{term['id']}/publications/1/send"))
    assert sent["sent_to_smarttt"] is True
    payload = smarttt.payloads[-1]
    assert payload["term_id"] is None and (payload["term_start"], payload["term_end"]) == ("2026-09-21", "2026-12-18")
    assert payload["based_on_version"] is None
    cosc = next(s for s in payload["sessions"] if s["unit_code"] == "COSC 103")
    assert cosc["department"] == {"smarttt_id": CS, "code": "DCOMP", "name": "Computer Science"}
    assert cosc["room"]["code"] in {"UTC 1", "STB 1", "STB 2", "STB 3"} and cosc["room"]["capacity"] > 0
    assert all(r["label"] and r["program_code"] for r in cosc["class_rows"])
    assert {r["department_smarttt_id"] for r in cosc["class_rows"]} <= {CS, EDU}
    assert ok(client.get(f"/api/terms/{term['id']}"))["smarttt_term_id"]
    db.expire_all()
    assert all(o.smarttt_unit_id for o in db.scalars(select(UnitOffering)))

    pdf = client.get(f"/api/drafts/{draft_id}/export.pdf?view=row")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    word = client.get(f"/api/drafts/{draft_id}/export.docx?view=row")
    assert word.status_code == 200 and word.content.startswith(b"PK")
    assert "wordprocessingml" in word.headers["content-type"]
    excel = client.get(f"/api/drafts/{draft_id}/export.xlsx?view=room")
    assert excel.status_code == 200 and excel.content.startswith(b"PK")
    assert client.get(f"/api/drafts/{draft_id}/export.txt").status_code == 404

    from docx import Document

    document = Document(BytesIO(word.content))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "SEPTEMBER-DECEMBER 2026 TEACHING TIMETABLE" in text
    assert "Prepared by Dr. Alex Mutegi Rwanda, Director Examination & Timetabling" in text
    assert "To be used from Monday 21/09/2026" in text
    table = document.tables[0]
    assert len(table.columns) == 61
    assert any("COSC 103" in c.text and "GR A" in c.text for row in table.rows for c in row.cells)


def test_attendance_reads_terms_and_published_times(client, smarttt):
    sign_in(client)
    term = enter_term(client)
    no_key = client.get("/api/integrations/attendance/terms")
    assert no_key.status_code == 401
    headers = {"X-API-Key": "test-attendance-key"}
    terms = ok(client.get("/api/integrations/attendance/terms", headers=headers))
    assert terms[0]["id"] == term["id"] and terms[0]["published_version"] is None

    before = ok(client.get(f"/api/integrations/attendance/terms/{term['id']}/times", headers=headers))
    assert before == {"published_version": None, "classes": {}}  # nothing published yet

    draft_id = ok(client.post("/api/runs", json={"term_id": term["id"], "time_limit_s": 20}))["draft_id"]
    ok(client.post(f"/api/drafts/{draft_id}/publish", json={}))
    after = ok(client.get(f"/api/integrations/attendance/terms/{term['id']}/times", headers=headers))
    assert after["published_version"] == 1
    assert set(after["classes"]) == {"COSC100-MAIN", "COSC103-GRA", "COSC103-GRB", "COSC103-GRC"}
    meeting = after["classes"]["COSC103-GRB"][0]
    assert meeting["room"] in {"UTC 1", "STB 1", "STB 2", "STB 3"}
    assert int(meeting["end_time"][:2]) - int(meeting["start_time"][:2]) == 3


def test_deleting_a_class_row_deletes_the_units_only_it_takes(client, smarttt):
    sign_in(client)
    term = enter_term(client)
    draft_id = ok(client.post("/api/runs", json={"term_id": term["id"], "time_limit_s": 20}))["draft_id"]
    ok(client.post(f"/api/drafts/{draft_id}/publish", json={}))
    rows = {r["label"]: r["id"] for r in ok(client.get(f"/api/terms/{term['id']}/class-rows"))}

    # COSC 100 is taken by BSC.COMP SCI alone, so it goes with the row and out of the draft.
    gone = ok(client.delete(f"/api/class-rows/{rows['BSC.COMP SCI Y1S1']}"))
    assert gone == {"units_deleted": ["COSC 100"], "units_kept": []}
    # COSC 103 is shared: BED.MATH/BIO leaves it and its GR C goes; GR A and GR B stay.
    left = ok(client.delete(f"/api/class-rows/{rows['BED.MATH/BIO Y1S1']}"))
    assert left == {"units_deleted": [], "units_kept": ["COSC 103"]}

    units = ok(client.get(f"/api/terms/{term['id']}/offerings"))
    assert [u["unit_code"] for u in units] == ["COSC 103"]
    assert [g["label"] for g in units[0]["groups"]] == ["GR A", "GR B"]
    assert "BED.MATH/BIO Y1S1" not in units[0]["class_rows"]
    # The approved version still shows what was approved.
    grid = ok(client.get(f"/api/drafts/{draft_id}?view=row"))
    assert grid["as_approved"] is True
    assert {c["unit_code"] for e in grid["entities"] for c in e["cells"]} == {"COSC 100", "COSC 103"}


def test_deleting_a_unit_takes_it_out_of_the_drafts(client, smarttt):
    sign_in(client)
    term = enter_term(client)
    draft_id = ok(client.post("/api/runs", json={"term_id": term["id"], "time_limit_s": 20}))["draft_id"]
    cosc103 = next(u for u in ok(client.get(f"/api/terms/{term['id']}/offerings")) if u["unit_code"] == "COSC 103")
    ok(client.delete(f"/api/offerings/{cosc103['id']}"))
    grid = ok(client.get(f"/api/drafts/{draft_id}?view=row"))
    assert {c["unit_code"] for e in grid["entities"] for c in e["cells"]} == {"COSC 100"}
    assert len(ok(client.get(f"/api/terms/{term['id']}/class-rows"))) == 4  # the class rows stay


def test_deleting_what_is_in_use_is_refused(client, smarttt):
    sign_in(client)
    term = enter_term(client)

    departments = {d["code"]: d for d in ok(client.get("/api/departments"))["departments"]}
    assert client.delete(f"/api/departments/{departments['DCOMP']['id']}").status_code == 409

    draft_id = ok(client.post("/api/runs", json={"term_id": term["id"], "time_limit_s": 20}))["draft_id"]
    grid = ok(client.get(f"/api/drafts/{draft_id}?view=room"))
    used_room = grid["entities"][0]["id"]
    assert client.delete(f"/api/rooms/{used_room}").status_code == 409
    closed = ok(client.patch(f"/api/rooms/{used_room}", json={"active": False}))
    assert closed["active"] is False


def test_duplicates_are_refused_and_bulk_updates(client, smarttt):
    sign_in(client)
    term = enter_term(client)
    dup = client.post("/api/rooms", json={"code": "utc 1", "capacity": 50})
    assert dup.status_code == 409 and dup.json()["error"]["code"] == "CODE_TAKEN"
    again = ok(client.post("/api/rooms/bulk", json={"rows": [{"code": "UTC 1", "capacity": "250"}]}))
    assert again == {"created": 0, "updated": 1, "errors": [], "error_count": 0}
    assert next(r for r in ok(client.get("/api/rooms")) if r["code"] == "UTC 1")["capacity"] == 250

    departments = {d["code"]: d for d in ok(client.get("/api/departments"))["departments"]}
    dup_unit = client.post(
        f"/api/terms/{term['id']}/offerings",
        json={"unit_code": "cosc100", "unit_name": "Again", "department_id": departments["DCOMP"]["id"]},
    )
    assert dup_unit.status_code == 409

    bulk_units = ok(
        client.post(
            f"/api/terms/{term['id']}/offerings/bulk",
            json={
                "rows": [
                    {
                        "unit_code": "COSC 104",
                        "unit_title": "Web Programming",
                        "department": "DCOMP",
                        "class_rows": "BSC.COMP SCI Y1S1",
                        "room_type": "Lab",
                    },
                    {
                        "unit_code": "COSC 105",
                        "unit_title": "Support",
                        "class_rows": "BSC.COMP SCI Y1S1",
                    },
                    {
                        "unit_code": "COSC 106",
                        "unit_title": "Install",
                        "department": "DCOMP",
                        "class_rows": "NO SUCH ROW Y1S1",
                    },
                ]
            },
        )
    )
    assert bulk_units["created"] == 1 and bulk_units["error_count"] == 2
    assert "department" in bulk_units["errors"][0]["message"]
    assert "class rows not found" in bulk_units["errors"][1]["message"]


def test_copy_a_term_into_the_next(client, smarttt):
    sign_in(client)
    first = enter_term(client)
    second = ok(
        client.post("/api/terms", json={"academic_year": "2026/2027", "semester": 2, "copy_from_term_id": first["id"]})
    )
    offerings = {o["unit_code"]: o for o in ok(client.get(f"/api/terms/{second['id']}/offerings"))}
    assert set(offerings) == {"COSC 100", "COSC 103"}
    assert sorted(g["label"] for g in offerings["COSC 103"]["groups"]) == ["GR A", "GR B", "GR C"]
    assert len(ok(client.get(f"/api/terms/{second['id']}/class-rows"))) == 4
    twice = ok(client.post(f"/api/terms/{second['id']}/copy", json={"source_term_id": first["id"]}))
    assert twice == {"class_rows_copied": 0, "units_copied": 0}
    clash = client.post("/api/terms", json={"academic_year": "2026/2027", "semester": 2})
    assert clash.status_code == 409


def test_departments_come_from_attendance_and_are_typed_by_name(client, smarttt, attendance):
    sign_in(client)
    term = ok(client.post("/api/terms", json={"academic_year": "2026/2027", "semester": 1, "name": "Sep-Dec 2026"}))
    cs = departments(client, {"Computer Science": {"code": "DCOMP"}})["Computer Science"]

    # An existing department is found by name (case, spacing and "Department of" ignored) or by code.
    row = ok(
        client.post(
            f"/api/terms/{term['id']}/class-rows",
            json={
                "program_code": "BSC.COMP SCI",
                "study_year": 1,
                "study_semester": 1,
                "department": "department of  computer science",
            },
        )
    )
    assert row["department_id"] == cs["id"]
    unit = ok(
        client.post(
            f"/api/terms/{term['id']}/offerings",
            json={"unit_code": "COSC 100", "unit_name": "Intro", "department": "dcomp", "class_row_ids": [row["id"]]},
        )
    )
    assert unit["department_id"] == cs["id"]

    # A department not in Attendance can't be typed into being here.
    refused = client.patch(f"/api/offerings/{unit['id']}", json={"department": "Mathematics and Statistics"})
    assert refused.status_code == 422 and refused.json()["error"]["code"] == "UNKNOWN_DEPARTMENT"
    assert "Smart Attendance" in refused.json()["error"]["message"]
    assert client.post("/api/departments", json={"code": "MS", "name": "Maths"}).status_code == 405

    # Once an administrator adds it in Attendance, it is read in and can be used. A department
    # Attendance has no code for (made before codes existed) gets one made here.
    attendance.add("att-math", "Department of Mathematics and Statistics")
    synced = ok(client.post("/api/departments/sync"))
    assert synced["report"]["created"] == 1
    maths = next(d for d in synced["departments"] if d["name"] == "Department of Mathematics and Statistics")
    assert maths["code"] == "MS" and maths["id"] >= 100_000 and maths["smarttt_department_id"] is None
    assert maths["code_from_attendance"] is False and "school" not in maths
    unit = ok(client.patch(f"/api/offerings/{unit['id']}", json={"department": "Mathematics and Statistics"}))
    assert unit["department_id"] == maths["id"]

    # The name and the code are Attendance's: changed there, they change here, and the code can't be changed here.
    assert cs["code_from_attendance"] is True
    attendance.departments[0]["name"] = "Computer Science and ICT"
    attendance.departments[0]["code"] = "DCSICT"
    renamed = {d["id"]: d for d in ok(client.post("/api/departments/sync"))["departments"]}[cs["id"]]
    assert (renamed["name"], renamed["code"]) == ("Computer Science and ICT", "DCSICT")
    assert client.patch(f"/api/departments/{cs['id']}", json={"name": "Other"}).status_code == 422
    refused_code = client.patch(f"/api/departments/{cs['id']}", json={"code": "OTHER"})
    assert refused_code.status_code == 409 and refused_code.json()["error"]["code"] == "CODE_FROM_ATTENDANCE"
    ok(client.patch(f"/api/departments/{cs['id']}", json={"lecturer_count": 4}))
    # Attendance's code for the new department replaces the one made here.
    attendance.departments[2]["code"] = "DMATHS"
    maths = {d["id"]: d for d in ok(client.post("/api/departments/sync"))["departments"]}[maths["id"]]
    assert (maths["code"], maths["code_from_attendance"]) == ("DMATHS", True)
    assert client.delete(f"/api/departments/{cs['id']}").json()["error"]["code"] == "FROM_ATTENDANCE"

    # A unit still needs a department, and a refused save adds none.
    response = client.post(f"/api/terms/{term['id']}/offerings", json={"unit_code": "COSC 101", "unit_name": "X"})
    assert response.status_code == 422
    assert len(ok(client.get("/api/departments"))["departments"]) == 3


def test_attendance_down_keeps_the_departments_as_last_read(client, smarttt, attendance):
    sign_in(client)
    assert {d["name"] for d in ok(client.get("/api/departments"))["departments"]} == {"Computer Science", "Education"}
    attendance.down = True
    listed = ok(client.get("/api/departments"))
    assert listed["attendance_connected"] is True and len(listed["departments"]) == 2
    forced = client.post("/api/departments/sync")
    assert forced.status_code == 503 and forced.json()["error"]["code"] == "ATTENDANCE_UNAVAILABLE"


def test_naming_a_group_replaces_the_whole_class(client, smarttt):
    """COSC 100 starts as one whole class; once split into GR A/GR B it no longer appears on its own."""
    sign_in(client)
    term = enter_term(client)
    draft_id = ok(client.post("/api/runs", json={"term_id": term["id"], "time_limit_s": 20}))["draft_id"]
    ok(client.post(f"/api/drafts/{draft_id}/publish", json={}))
    cosc100 = next(o for o in ok(client.get(f"/api/terms/{term['id']}/offerings")) if o["unit_code"] == "COSC 100")
    row_id = cosc100["groups"][0]["members"][0]["class_row_id"]

    for label, students in (("GR A", 45), ("GR B", 45)):
        ok(
            client.post(
                f"/api/offerings/{cosc100['id']}/groups",
                json={"label": label, "members": [{"class_row_id": row_id, "students": students}]},
            )
        )
    cosc100 = next(o for o in ok(client.get(f"/api/terms/{term['id']}/offerings")) if o["unit_code"] == "COSC 100")
    assert [g["label"] for g in cosc100["groups"]] == ["GR A", "GR B"]

    # The approved version still reads as approved; the next draft needs the groups (it is out of date).
    grid = ok(client.get(f"/api/drafts/{draft_id}?view=row"))
    assert grid["as_approved"] is True
    assert any(c["unit_code"] == "COSC 100" and c["group"] == "" for e in grid["entities"] for c in e["cells"])
    copy = ok(client.post(f"/api/drafts/{draft_id}/copy"))
    assert set(ok(client.get(f"/api/drafts/{copy['id']}?view=row"))["draft"]["stale"]["missing"]) >= {
        "COSC100-GRA-1",
        "COSC100-GRB-1",
    }

    # The whole class can't come back beside its groups.
    again = client.post(
        f"/api/offerings/{cosc100['id']}/groups",
        json={"label": "", "members": [{"class_row_id": row_id, "students": 90}]},
    )
    assert again.status_code == 409 and again.json()["error"]["code"] == "UNIT_SPLIT"


def test_a_file_naming_something_twice_is_refused_on_the_second_row(client, smarttt):
    sign_in(client)
    term = enter_term(client)
    rooms = ok(
        client.post(
            "/api/rooms/bulk",
            json={"rows": [{"code": "LAB 7", "capacity": "40"}, {"code": "lab7", "capacity": "400"}]},
        )
    )
    assert rooms["created"] == 1
    assert rooms["errors"] == [{"line": 2, "message": "Lecture hall lab7 is also on row 1"}]
    assert next(r for r in ok(client.get("/api/rooms")) if r["code"] == "LAB 7")["capacity"] == 40

    counts = ok(
        client.post(
            "/api/departments/bulk",
            json={
                "rows": [
                    {"department": "Department of Computer Science", "lecturers": "8"},
                    {"department": "dcomp", "lecturers": "9"},
                    {"department": "Department of Chemistry", "lecturers": "3"},
                ]
            },
        )
    )
    assert counts["updated"] == 1 and counts["error_count"] == 2
    assert counts["errors"][0]["message"] == "Department DCOMP is also on row 1"
    assert "Smart Attendance" in counts["errors"][1]["message"]

    unit = {"unit_title": "Data Structures", "department": "DCOMP", "class_rows": "BSC.COMP SCI Y1S1"}
    units = ok(
        client.post(
            f"/api/terms/{term['id']}/offerings/bulk",
            json={"rows": [{"unit_code": "COSC 210", **unit}, {"unit_code": "COSC210", **unit}]},
        )
    )
    assert units["created"] == 1
    assert units["errors"][0]["message"] == "Unit COSC210 is also on row 1"


def test_cosc_units_go_to_a_lab_when_one_seats_the_group(client, smarttt):
    sign_in(client)
    term = enter_term(client)  # LAB 1 seats 40; COSC 100 has 90 students, COSC 103's groups 100+
    departments = {d["code"]: d for d in ok(client.get("/api/departments"))["departments"]}
    ok(
        client.post(
            f"/api/terms/{term['id']}/class-rows/bulk",
            json={
                "rows": [{"programme": "BSC.IT", "year": "1", "semester": "1", "department": "DCOMP", "students": "35"}]
            },
        )
    )
    labels = {r["label"]: r["id"] for r in ok(client.get(f"/api/terms/{term['id']}/class-rows"))}
    for code, name in (("COSC 110", "Programming"), ("MATH 110", "Calculus")):
        ok(
            client.post(
                f"/api/terms/{term['id']}/offerings",
                json={
                    "unit_code": code,
                    "unit_name": name,
                    "department_id": departments["DCOMP"]["id"],
                    "class_row_ids": [labels["BSC.IT Y1S1"]],
                    "room_type": "lecture_hall",
                },
            )
        )
    run = ok(client.post("/api/runs", json={"term_id": term["id"], "time_limit_s": 20}))
    assert run["status"] == "succeeded", run["failure"]
    grid = ok(client.get(f"/api/drafts/{run['draft_id']}?view=row"))
    assert grid["draft"]["violation_count"] == 0
    rooms = {c["unit_code"]: c["room"] for e in grid["entities"] for c in e["cells"]}
    assert rooms["COSC 110"] == "LAB 1"  # 35 students: the lab, though the unit says lecture hall
    assert rooms["MATH 110"] != "LAB 1"  # not a computing unit
    assert rooms["COSC 100"] != "LAB 1"  # 90 students: too many for the lab


def test_room_type_for_computing_units():
    from app.services.problem import room_type_for

    largest = {"lab": 40, "lecture_hall": 200}
    assert room_type_for("COSC 110", "lecture_hall", 35, largest) == "lab"
    assert room_type_for("cosc110", "seminar", 40, largest) == "lab"
    assert room_type_for("COSC 100", "lecture_hall", 90, largest) == "lecture_hall"
    assert room_type_for("COSC 100", "lab", 90, largest) == "lecture_hall"  # too big for every lab
    assert room_type_for("MATH 110", "lecture_hall", 35, largest) == "lecture_hall"
    assert room_type_for("COSC 110", "lecture_hall", 35, {"lecture_hall": 200}) == "lecture_hall"  # no labs
