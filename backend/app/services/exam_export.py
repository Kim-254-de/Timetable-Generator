"""
Exam downloads, Word and PDF:

- the exam timetable in the Directorate's layout: THARAKA UNIVERSITY,
  DIRECTORATE OF EXAMINATION & TIMETABLING, the title line, a DATE/DAY x
  session grid of `paper -> rooms`, practicals on their dates, "Prepared by";
- room sheets: per date and session, each room's papers, class rows,
  students and invigilator;
- the invigilation roster per department.
"""

from __future__ import annotations

from collections import defaultdict
from datetime import date
from io import BytesIO
from typing import Any
from xml.sax.saxutils import escape

from app.models import Term
from app.services.exam_drafts import MAIN, room_sheets

UNIVERSITY = "THARAKA UNIVERSITY"
DIRECTORATE = "DIRECTORATE OF EXAMINATION & TIMETABLING"


def _title(term: Term, version: int | None) -> list[str]:
    name = (term.name or f"{term.academic_year} semester {term.semester}").upper()
    line = f"END OF {name} EXAMINATION TIMETABLE"
    return [UNIVERSITY, DIRECTORATE, line + (f" - VERSION {version}" if version else " - DRAFT")]


def _prepared(term: Term, printed_by: str) -> str:
    return f"Prepared by: {term.prepared_by or printed_by}, {term.prepared_by_title}"


def _session_label(s: dict[str, str]) -> str:
    return f"{s['start']}-{s['end']}"


def timetable_table(rec: dict[str, Any]) -> tuple[list[str], list[list[str]]]:
    """The header and rows of the date x session grid; each cell lists `paper -> rooms`."""
    cells: dict[tuple[str, int], list[str]] = defaultdict(list)
    for p in rec["papers"]:
        if p["day"] is None:
            continue
        rooms = sorted({s["room"] for s in p["seating"]})
        cells[(p["day"], p["session"])].append(f"{p['codes']} → {', '.join(rooms)}")
    header = ["DATE/DAY", *[_session_label(s) for s in rec["sessions"]]]
    rows = []
    for d in rec["days"]:
        rows.append([d["label"], *["\n".join(sorted(cells[(d["date"], i)])) for i in range(len(rec["sessions"]))]])
    return header, rows


def practical_lines(rec: dict[str, Any]) -> list[str]:
    return [
        f"{date.fromisoformat(p['date']):%a %d/%m/%Y}  {p['start_time']}-{p['end_time']}  "
        f"{p['codes']} {p['title']} PRACTICALS"
        for p in rec["practicals"]
    ]


ROSTER_HEADERS = ["Paper", "Date", "Session", "Room", "Duty", "Invigilator"]
ROOM_HEADERS = ["Room", "Paper", "Class row", "Students", "Invigilator"]


def roster(rec: dict[str, Any]) -> dict[str, list[list[str]]]:
    """Per department: each paper's main invigilator (collects every room's scripts), then one per extra room."""
    out: dict[str, list[list[str]]] = defaultdict(list)
    days = {d["date"]: d["label"] for d in rec["days"]}
    for p in rec["papers"]:
        if p["day"] is None:
            continue
        when = [days.get(p["day"], p["day"]), _session_label(rec["sessions"][p["session"]])]
        main = p["invigilators"].get(MAIN) or {}
        out[p["department"]].append([p["codes"], *when, p["main_room"] or "", "Main", main.get("name") or "Not named"])
        for room in p["extra_rooms"]:
            extra = p["invigilators"].get(room) or {}
            out[p["department"]].append([p["codes"], *when, room, "Extra", extra.get("name") or "Not named"])
    return dict(sorted(out.items()))


def _room_rows(slot: dict[str, Any]) -> list[list[str]]:
    rows = []
    for room in slot["rooms"]:
        for paper in room["papers"]:
            who = paper["invigilator"] or "Not named"
            who += " (main)" if paper["main"] else f" (extra; main: {paper['main_invigilator'] or 'not named'})"
            for row in paper["rows"]:
                odel = " (ODEL)" if row["odel"] else ""
                rows.append([room["room"], paper["codes"], row["row"] + odel, str(row["students"]), who])
    return rows


# -- Word ---------------------------------------------------------------------------------------


def _docx_heading(doc: Any, lines: list[str]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    for line in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        run.bold = True
        run.font.size = Pt(12)


def _docx_table(doc: Any, header: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows) + 1, cols=len(header))
    table.style = "Table Grid"
    for r, values in enumerate([header, *rows]):
        for c, value in enumerate(values):
            cell = table.rows[r].cells[c]
            cell.text = value
            if r == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def _new_docx(landscape: bool) -> Any:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt

    doc = Document()
    section = doc.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = section.right_margin = Cm(1.2)
    section.top_margin = section.bottom_margin = Cm(1.2)
    doc.styles["Normal"].font.size = Pt(8)
    return doc


def _save(doc: Any) -> bytes:
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def timetable_docx(term: Term, rec: dict[str, Any], version: int | None, printed_by: str) -> bytes:
    doc = _new_docx(landscape=True)
    _docx_heading(doc, _title(term, version))
    header, rows = timetable_table(rec)
    _docx_table(doc, header, rows)
    lines = practical_lines(rec)
    if lines:
        doc.add_paragraph().add_run("PRACTICALS").bold = True
        for line in lines:
            doc.add_paragraph(line)
    doc.add_paragraph()
    doc.add_paragraph(_prepared(term, printed_by))
    return _save(doc)


def rooms_docx(term: Term, rec: dict[str, Any], version: int | None, printed_by: str) -> bytes:
    doc = _new_docx(landscape=False)
    _docx_heading(doc, [*_title(term, version), "ROOM SHEETS"])
    days = {d["date"]: d["label"] for d in rec["days"]}
    for slot in room_sheets(rec):
        session = rec["sessions"][slot["session"]]
        doc.add_paragraph().add_run(f"{days.get(slot['day'], slot['day'])}  {_session_label(session)}").bold = True
        rows = _room_rows(slot)
        _docx_table(doc, ROOM_HEADERS, rows)
    doc.add_paragraph(_prepared(term, printed_by))
    return _save(doc)


def roster_docx(term: Term, rec: dict[str, Any], version: int | None, printed_by: str) -> bytes:
    doc = _new_docx(landscape=False)
    _docx_heading(doc, [*_title(term, version), "INVIGILATION ROSTER"])
    for department, rows in roster(rec).items():
        doc.add_paragraph().add_run(department).bold = True
        _docx_table(doc, ROSTER_HEADERS, rows)
    doc.add_paragraph(_prepared(term, printed_by))
    return _save(doc)


# -- PDF ----------------------------------------------------------------------------------------


def _pdf(story_for: Any, landscape_page: bool, title: str) -> bytes:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate

    out = BytesIO()
    SimpleDocTemplate(
        out,
        pagesize=landscape(A4) if landscape_page else A4,
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
        title=title,
    ).build(story_for())
    return out.getvalue()


def _styles() -> dict[str, Any]:
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet

    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("title", parent=base["Title"], fontSize=11, leading=14, spaceAfter=0),
        "cell": ParagraphStyle("cell", parent=base["BodyText"], fontSize=7, leading=8.5),
        "heading": ParagraphStyle("heading", parent=base["Heading4"], spaceBefore=6, spaceAfter=3),
        "body": base["BodyText"],
    }


def _pdf_table(header: list[str], rows: list[list[str]], widths: list[float] | None = None) -> Any:
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Table, TableStyle

    st = _styles()
    data = [[Paragraph(f"<b>{escape(h)}</b>", st["cell"]) for h in header]] + [
        [Paragraph(escape(v).replace("\n", "<br/>"), st["cell"]) for v in row] for row in rows
    ]
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    return table


def timetable_pdf(term: Term, rec: dict[str, Any], version: int | None, printed_by: str) -> bytes:
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Spacer

    st = _styles()
    title = _title(term, version)

    def story() -> list[Any]:
        header, rows = timetable_table(rec)
        first = 28 * mm
        rest = (277 * mm - first) / max(1, len(header) - 1)
        out: list[Any] = [Paragraph(escape(line), st["title"]) for line in title]
        out += [Spacer(1, 4 * mm), _pdf_table(header, rows, [first] + [rest] * (len(header) - 1))]
        lines = practical_lines(rec)
        if lines:
            out += [Paragraph("PRACTICALS", st["heading"]), *[Paragraph(escape(line), st["cell"]) for line in lines]]
        out += [Spacer(1, 6 * mm), Paragraph(escape(_prepared(term, printed_by)), st["body"])]
        return out

    return _pdf(story, True, title[-1])


def rooms_pdf(term: Term, rec: dict[str, Any], version: int | None, printed_by: str) -> bytes:
    from reportlab.platypus import Paragraph

    st = _styles()
    title = [*_title(term, version), "ROOM SHEETS"]
    days = {d["date"]: d["label"] for d in rec["days"]}

    def story() -> list[Any]:
        out: list[Any] = [Paragraph(escape(line), st["title"]) for line in title]
        for slot in room_sheets(rec):
            session = rec["sessions"][slot["session"]]
            out.append(
                Paragraph(escape(f"{days.get(slot['day'], slot['day'])}  {_session_label(session)}"), st["heading"])
            )
            rows = _room_rows(slot)
            out.append(_pdf_table(ROOM_HEADERS, rows))
        out.append(Paragraph(escape(_prepared(term, printed_by)), st["body"]))
        return out

    return _pdf(story, False, title[-1])


def roster_pdf(term: Term, rec: dict[str, Any], version: int | None, printed_by: str) -> bytes:
    from reportlab.platypus import Paragraph

    st = _styles()
    title = [*_title(term, version), "INVIGILATION ROSTER"]

    def story() -> list[Any]:
        out: list[Any] = [Paragraph(escape(line), st["title"]) for line in title]
        for department, rows in roster(rec).items():
            out.append(Paragraph(escape(department), st["heading"]))
            out.append(_pdf_table(ROSTER_HEADERS, rows))
        out.append(Paragraph(escape(_prepared(term, printed_by)), st["body"]))
        return out

    return _pdf(story, False, title[-1])


BUILDERS = {
    ("timetable", "docx"): timetable_docx,
    ("timetable", "pdf"): timetable_pdf,
    ("rooms", "docx"): rooms_docx,
    ("rooms", "pdf"): rooms_pdf,
    ("roster", "docx"): roster_docx,
    ("roster", "pdf"): roster_pdf,
}
MEDIA = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
