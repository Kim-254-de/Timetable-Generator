"""
A draft as a PDF, a Word document or an Excel workbook, in the layout of
today's teaching timetable: the Directorate's header, one row per class row
(or room, or department), a Monday-Friday x 07:00-19:00 grid, cells
"UNIT / GR X / ROOM" spanning the class's hours, the version line and the
"Prepared by ... Director Examination & Timetabling" footer.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models import Term

DAY_NAMES = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday")
TERM_MONTHS = {1: "SEPTEMBER-DECEMBER", 2: "JANUARY-APRIL", 3: "MAY-AUGUST"}


def term_title(term: Term) -> str:
    if term.name:
        return f"{term.name.upper()} TEACHING TIMETABLE"
    first, _, second = term.academic_year.partition("/")
    year = first if term.semester == 1 else (second or first)
    return f"{TERM_MONTHS.get(term.semester, f'SEMESTER {term.semester}')} {year} TEACHING TIMETABLE"


def prepared_by(term: Term, printed_by: str) -> str:
    title = term.prepared_by_title or "Director Examination & Timetabling"
    return f"Prepared by {term.prepared_by or printed_by}, {title}"


def version_line(term: Term, grid: dict[str, Any]) -> str:
    draft = grid["draft"]
    parts = [term.version_label or f"{draft['name']} ({draft['status']})"]
    if term.use_from:
        parts.append(f"To be used from {term.use_from.strftime('%A %d/%m/%Y')}")
    return " - ".join(parts)


def cell_lines(grid: dict[str, Any], cell: dict[str, Any]) -> list[str]:
    fields = (
        ("unit_code", "group", "lecturer") if grid["view"] == "room" else ("unit_code", "group", "room", "lecturer")
    )
    return [str(cell[f]) for f in fields if cell.get(f)]


def stacked_lines(entity: dict[str, Any]) -> list[list[dict[str, Any]]]:
    """A row can hold parallel classes at the same hour (rooms, departments): stack them in extra lines."""
    lines: list[list[dict[str, Any]]] = []
    for cell in entity["cells"]:
        for line in lines:
            if all(
                cell["start_slot"] >= c["start_slot"] + c["length"]
                or c["start_slot"] >= cell["start_slot"] + cell["length"]
                for c in line
            ):
                line.append(cell)
                break
        else:
            lines.append([cell])
    return lines or [[]]


def build_pdf(term: Term, grid: dict[str, Any], printed_by: str) -> bytes:
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=5.5, leading=6.5, alignment=1)
    label_style = ParagraphStyle("label", parent=styles["BodyText"], fontSize=6.5, leading=7.5)
    hours = grid["hours"]
    first_hour = grid["first_hour"]

    header_days: list[Any] = [""]
    header_hours: list[Any] = [""]
    for day in DAY_NAMES:
        header_days += [day] + [""] * (hours - 1)
        header_hours += [f"{h}-{h + 1}" for h in range(first_hour, first_hour + hours)]
    data: list[list[Any]] = [header_days, header_hours]
    style: list[Any] = [
        ("FONT", (0, 0), (-1, 1), "Helvetica-Bold", 6),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.black),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (1, 0), (-1, 1), "CENTER"),
        ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#e5e7eb")),
    ]
    for d in range(len(DAY_NAMES)):
        style.append(("SPAN", (1 + d * hours, 0), (d * hours + hours, 0)))

    for entity in grid["entities"]:
        for n, line in enumerate(stacked_lines(entity)):
            row: list[Any] = [Paragraph(entity["label"], label_style) if n == 0 else ""] + [""] * (5 * hours)
            r = len(data)
            for cell in line:
                text = "<br/>".join(cell_lines(grid, cell))
                row[1 + cell["start_slot"]] = Paragraph(text, cell_style)
                style.append(("SPAN", (1 + cell["start_slot"], r), (cell["start_slot"] + cell["length"], r)))
                if cell["pinned"]:
                    style.append(
                        (
                            "BACKGROUND",
                            (1 + cell["start_slot"], r),
                            (cell["start_slot"] + cell["length"], r),
                            colors.HexColor("#fef3c7"),
                        )
                    )
            data.append(row)

    width = landscape(A3)[0] - 20 * mm
    label_width = 32 * mm
    col = (width - label_width) / (5 * hours)
    table = Table(data, colWidths=[label_width] + [col] * (5 * hours), repeatRows=2)
    table.setStyle(TableStyle(style))

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A3),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=14 * mm,
        title=term_title(term),
        author=printed_by,
    )
    footer_left = prepared_by(term, printed_by)
    footer_right = version_line(term, grid)

    def footer(canvas: Any, document: Any) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawString(10 * mm, 7 * mm, footer_left)
        canvas.drawRightString(landscape(A3)[0] - 10 * mm, 7 * mm, f"{footer_right} - Page {document.page}")
        canvas.restoreState()

    story = [
        Paragraph("DIRECTORATE OF EXAMINATIONS AND TIMETABLING", styles["Title"]),
        Paragraph(term_title(term), styles["Heading2"]),
        Spacer(1, 3 * mm),
        table,
    ]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def build_docx(term: Term, grid: dict[str, Any], printed_by: str) -> bytes:
    """
    The same timetable as an editable Word document: an A3 landscape page and
    a real table (61 columns: the row label and Monday-Friday x 12 hours),
    each class a merged cell over its hours, so the Directorate can still
    adjust the printout by hand.
    """
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    hours = grid["hours"]
    first_hour = grid["first_hour"]
    columns = 1 + 5 * hours

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Mm(420), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Mm(10))

    def centred(text: str, size: float, bold: bool = True) -> None:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)

    centred("THARAKA UNIVERSITY", 16)
    centred("DIRECTORATE OF EXAMINATIONS AND TIMETABLING", 13)
    centred(term_title(term), 12)

    def write(cell: Any, lines: list[str], size: float = 5.5, bold: bool = False) -> None:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        for i, line in enumerate(lines):
            run = paragraph.add_run(line)
            run.font.size = Pt(size)
            run.bold = bold
            if i < len(lines) - 1:
                run.add_break()

    def shade(cell: Any, fill: str) -> None:
        properties = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), fill)
        properties.append(shading)

    body_rows = [(entity, line, n) for entity in grid["entities"] for n, line in enumerate(stacked_lines(entity))]
    table = document.add_table(rows=2 + len(body_rows), cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Cells are fetched one row at a time: table.cell() walks the whole table
    # on every call, which is far too slow for a 400-row timetable.
    label_width = Mm(32)
    hour_width = int((Mm(400) - label_width) / (5 * hours))
    rows = list(table.rows)
    # Word reads the grid columns, LibreOffice and Google Docs the cells: set both.
    for c, column in enumerate(table.columns):
        column.width = label_width if c == 0 else hour_width
    for row in rows:
        cells = row.cells
        cells[0].width = label_width
        for c in range(1, columns):
            cells[c].width = hour_width

    day_cells, hour_cells = rows[0].cells, rows[1].cells
    for d, day in enumerate(DAY_NAMES):
        merged = day_cells[1 + d * hours].merge(day_cells[d * hours + hours])
        write(merged, [day], 7, bold=True)
        shade(merged, "E5E7EB")
        for h in range(hours):
            hour_cell = hour_cells[1 + d * hours + h]
            write(hour_cell, [f"{first_hour + h}-{first_hour + h + 1}"], 5, bold=True)
            shade(hour_cell, "E5E7EB")
    shade(day_cells[0], "E5E7EB")
    shade(hour_cells[0], "E5E7EB")
    # Repeat the two header rows on every page.
    for header in table.rows[:2]:
        header_properties = header._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        header_properties.append(repeat)

    for r, (entity, line, n) in enumerate(body_rows, start=2):
        cells = rows[r].cells
        if n == 0:
            write(cells[0], [entity["label"]], 6.5, bold=True)
        for cell in line:
            first = 1 + cell["start_slot"]
            last = cell["start_slot"] + cell["length"]
            target = cells[first] if first == last else cells[first].merge(cells[last])
            write(target, cell_lines(grid, cell))
            if cell.get("pinned"):
                shade(target, "FEF3C7")

    document.add_paragraph()
    document.add_paragraph(version_line(term, grid)).runs[0].font.size = Pt(8)
    document.add_paragraph(prepared_by(term, printed_by)).runs[0].font.size = Pt(9)
    signature = document.add_paragraph("Signature: ______________________    Date: ______________")
    signature.runs[0].font.size = Pt(9)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_xlsx(term: Term, grid: dict[str, Any], printed_by: str) -> bytes:
    """One line per class, for checking and sorting in Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Timetable"
    sheet.append([term_title(term)])
    sheet["A1"].font = Font(bold=True, size=13)
    sheet.append([version_line(term, grid)])
    sheet.append([])
    heading = {"row": "Class row", "room": "Lecture hall", "department": "Department"}[grid["view"]]
    sheet.append([heading, "Unit", "Group", "Lecture hall", "Day", "Start", "End", "Hours", "Students"])
    for cell in sheet[4]:
        cell.font = Font(bold=True)
    day_names = dict(zip(("MON", "TUE", "WED", "THU", "FRI"), DAY_NAMES, strict=True))
    for entity in grid["entities"]:
        for c in entity["cells"]:
            sheet.append(
                [
                    entity["label"],
                    c["unit_code"],
                    c["group"] or "",
                    c["room"],
                    day_names.get(c["day"], c["day"]),
                    c["start_time"],
                    c["end_time"],
                    c["length"],
                    c.get("students"),
                ]
            )
    for column, width in zip("ABCDEFGHI", (30, 14, 10, 14, 12, 8, 8, 7, 9), strict=True):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A5"
    sheet.append([])
    sheet.append([prepared_by(term, printed_by)])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
